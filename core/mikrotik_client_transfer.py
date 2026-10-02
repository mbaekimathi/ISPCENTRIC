"""Bulk transfer of billing clients between MikroTik routers."""

from __future__ import annotations

import logging
import re
import threading
from datetime import datetime
from io import BytesIO
from typing import Any, Iterable

from django.core.cache import cache
from django.db.models import Count, Q
from django.http import HttpResponse
from django.utils import timezone

from billing.models import Customer
from billing.services import customer_needs_nas_provision
from core.models import MikroTikRouter
from core.subscription_sync import enqueue_customer_subscription_sync

logger = logging.getLogger(__name__)

_EXCEL_HEADERS = (
    "Full name",
    "Account number",
    "Phone",
    "Email",
    "Status",
    "Service",
    "PPPoE username",
    "Package",
    "Package start",
    "Package end",
    "MikroTik",
    "MikroTik host",
    "Location",
    "Building",
    "House number",
    "CPE IP",
)


def transfer_customer_to_router(customer: Customer, target_router: MikroTikRouter) -> dict[str, Any]:
    """
    Reassign one client to ``target_router`` and enqueue NAS provision when needed.

    Mirrors the single-client ``update_client_details`` router_changed path.
    """
    from core.mikrotik_connect import (
        provision_customer_pppoe,
        provision_static_client_dhcp_lease,
    )

    old_router_id = customer.router_id
    target_id = getattr(target_router, "pk", None)
    if old_router_id == target_id:
        return {"moved": False, "reason": "already_on_target"}

    customer.router = target_router
    customer.save(update_fields=["router"])

    org_id = getattr(customer, "organization_id", None)
    if org_id:
        cache.delete(f"client_cpe_router_data:{org_id}:{customer.pk}")
        cache.delete(f"client_cpe_wifi:{org_id}:{customer.pk}")

    sync_pppoe = (
        customer.service_type == Customer.ServiceType.PPPOE
        and bool(customer.router_id)
        and bool((customer.pppoe_username or "").strip())
    )
    static_dhcp_bind = (
        customer.service_type == Customer.ServiceType.STATIC
        and bool(customer.router_id)
        and bool((customer.cpe_ip or "").strip())
        and bool((customer.cpe_mac or "").strip())
    )

    if sync_pppoe:

        def _bg_provision(pk: int = customer.pk) -> None:
            from django.db import connection

            try:
                cust = Customer.objects.select_related(
                    "plan", "router", "organization"
                ).get(pk=pk)
                if cust.pppoe_username and cust.router_id:
                    provision_customer_pppoe(cust, ensure_stack=False)
            except Exception:
                logger.exception("PPPoE provision failed after transfer for customer %s", pk)
            finally:
                connection.close()

        threading.Thread(target=_bg_provision, daemon=True).start()

    if static_dhcp_bind:

        def _bg_bind(pk: int = customer.pk) -> None:
            from django.db import connection

            try:
                cust = Customer.objects.select_related("router", "organization").get(
                    pk=pk
                )
                provision_static_client_dhcp_lease(cust)
            except Exception:
                logger.exception("Static DHCP bind failed after transfer for customer %s", pk)
            finally:
                connection.close()

        threading.Thread(target=_bg_bind, daemon=True).start()

    if customer.router_id and (customer.pppoe_username or "").strip():
        enqueue_customer_subscription_sync(
            customer.pk,
            customer_needs_nas_provision(customer),
            wait_first=True,
            quick=True,
        )

    return {
        "moved": True,
        "from_router_id": old_router_id,
        "to_router_id": target_id,
        "sync_pppoe": sync_pppoe,
        "static_dhcp": static_dhcp_bind,
    }


def _excel_dt(value) -> datetime | str:
    if value is None:
        return ""
    if timezone.is_aware(value):
        value = timezone.localtime(value)
    return value.replace(tzinfo=None) if hasattr(value, "replace") else value


def _safe_excel_filename(org_name: str = "", *, selected: bool = False) -> str:
    stamp = timezone.localtime().strftime("%Y%m%d-%H%M")
    org_bit = re.sub(r"[^\w\-]+", "-", (org_name or "").strip())[:40].strip("-")
    scope = "selected" if selected else "clients"
    parts = ["ispcentric", scope, org_bit, stamp]
    return "-".join(p for p in parts if p) + ".xlsx"


def clients_excel_response(
    customers: Iterable[Customer],
    *,
    org_name: str = "",
    selected: bool = False,
) -> HttpResponse:
    """Build an .xlsx download for the given PPPoE clients (Excel-compatible)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    rows = list(customers)
    rows.sort(
        key=lambda c: (
            (getattr(getattr(c, "router", None), "name", None) or "zzz").lower(),
            (c.full_name or "").lower(),
            c.pk or 0,
        )
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Clients"
    ws.append(list(_EXCEL_HEADERS))

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="0D9488")
    header_align = Alignment(vertical="center", wrap_text=True)
    for col_idx, _ in enumerate(_EXCEL_HEADERS, start=1):
        cell = ws.cell(1, col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    for customer in rows:
        router = getattr(customer, "router", None)
        plan = getattr(customer, "plan", None)
        ws.append(
            [
                customer.full_name or "",
                customer.account_number or "",
                customer.phone or "",
                customer.email or "",
                customer.get_status_display() if hasattr(customer, "get_status_display") else (customer.status or ""),
                customer.get_service_type_display()
                if hasattr(customer, "get_service_type_display")
                else (customer.service_type or ""),
                customer.pppoe_username or "",
                getattr(plan, "name", "") or "",
                _excel_dt(getattr(customer, "package_start", None)),
                _excel_dt(getattr(customer, "package_end", None)),
                getattr(router, "name", "") or "No MikroTik assigned",
                getattr(router, "host", "") or "",
                customer.address or "",
                customer.building_name or "",
                customer.house_number or "",
                getattr(customer, "cpe_ip", "") or "",
            ]
        )

    widths = (22, 16, 14, 24, 12, 10, 16, 16, 18, 18, 18, 16, 24, 16, 12, 14)
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"

    buf = BytesIO()
    wb.save(buf)
    payload = buf.getvalue()
    filename = _safe_excel_filename(org_name, selected=selected)
    response = HttpResponse(
        payload,
        content_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    response["Cache-Control"] = "no-store"
    return response


def build_transfer_groups(org, *, focus_router_id: int | None = None) -> tuple[list, list, list]:
    """Return (routers, PPPoE customers, groups) for the transfer page.

    Hotspot clients are excluded — they are not provisioned per-router the same way.
    """
    pppoe = Customer.ServiceType.PPPOE
    routers = list(
        MikroTikRouter.objects.filter(organization=org)
        .annotate(
            customer_count=Count(
                "customers",
                filter=Q(customers__service_type=pppoe),
            )
        )
        .order_by("name")
    )
    customers = list(
        Customer.objects.filter(organization=org, service_type=pppoe)
        .select_related("router", "plan")
        .order_by("full_name", "id")
    )

    groups: list[dict[str, Any]] = []
    for router in routers:
        groups.append(
            {
                "key": f"r{router.pk}",
                "router": router,
                "router_id": router.pk,
                "title": router.name,
                "subtitle": router.host or "",
                "customers": [c for c in customers if c.router_id == router.pk],
                "focused": focus_router_id == router.pk,
            }
        )
    unassigned = [c for c in customers if not c.router_id]
    groups.append(
        {
            "key": "unassigned",
            "router": None,
            "router_id": None,
            "title": "No MikroTik assigned",
            "subtitle": "Clients not linked to a router yet",
            "customers": unassigned,
            "focused": False,
        }
    )
    if focus_router_id:
        groups.sort(key=lambda g: (0 if g.get("focused") else 1, g["title"].lower()))
    return routers, customers, groups
